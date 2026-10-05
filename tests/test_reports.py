import hashlib
from datetime import date, timedelta
from decimal import Decimal

import pytest

from apps.core.enums import Channel, MatchStatus, OtomaxCategory
from apps.core.normalize import extract_tokens, norm_ref, ref_core
from apps.ingest.models import BankMutation, ImportBatch, OtomaxEntry
from apps.recon.engine import run_match
from apps.recon.reports import get_daily_summary, get_range_summary

BD = date(2026, 9, 11)
BD_NEXT = BD + timedelta(days=1)


def _h(*parts) -> str:
    return hashlib.sha256("|".join(str(p) for p in parts).encode()).hexdigest()


def _batch(channel, book_date):
    return ImportBatch.objects.create(
        channel=channel, book_date=book_date, source_filename="t", file_hash=_h(channel, book_date)
    )


def _bank(desc, amount, channel=Channel.BRI, book_date=BD):
    return BankMutation.objects.create(
        import_batch=_batch(channel, book_date),
        channel=channel,
        book_date=book_date,
        description_raw=desc,
        ref_normalized=norm_ref(desc),
        ref_core=ref_core(desc),
        extracted_tokens=extract_tokens(desc),
        amount=Decimal(amount),
        row_hash=_h("b", desc, amount, book_date),
    )


def _otomax(desc, amount, channel=Channel.BRI, book_date=BD):
    return OtomaxEntry.objects.create(
        import_batch=_batch(Channel.OTOMAX, book_date),
        book_date=book_date,
        reseller_name_raw="X",
        amount=Decimal(amount),
        description_raw=desc,
        category=OtomaxCategory.TOPUP_TARTUN,
        channel_hint=channel,
        ref_normalized=norm_ref(desc),
        ref_core=ref_core(desc),
        extracted_tokens=extract_tokens(desc),
        row_hash=_h("o", desc, amount, book_date),
    )


def _otomax_other(
    desc, amount, category, channel_hint="", book_date=BD, match_status=MatchStatus.UNMATCHED
):
    return OtomaxEntry.objects.create(
        import_batch=_batch(Channel.OTOMAX, book_date),
        book_date=book_date,
        reseller_name_raw="X",
        amount=Decimal(amount),
        description_raw=desc,
        category=category,
        channel_hint=channel_hint,
        ref_normalized=norm_ref(desc),
        ref_core=ref_core(desc),
        match_status=match_status,
        row_hash=_h("o-other", desc, amount, book_date, category),
    )


@pytest.mark.django_db
def test_selisih_ignores_matches_recorded_on_a_later_otomax_book_date():
    """Reproduksi kasus nyata: 44 mutasi bank BRI di tanggal BD sudah matched ke entry
    Otomax yang diinput operator H+1 (lazim buat QRIS), plus 1 mutasi bank yang memang
    tidak ada pasangannya sama sekali. 'Selisih' harus cuma menghitung sisa PR riil
    (Rp 20.500), bukan total_bank - total_otomax yang meledak jadi besar gara-gara
    Otomax-nya "pindah tanggal" ke BD_NEXT."""
    _bank("REF-COCOK-001", "203670409", book_date=BD)
    _otomax("REF-COCOK-001", "203670409", book_date=BD_NEXT)
    _bank("REF-TANPA-PASANGAN", "20500", book_date=BD)

    run_match(BD)

    summary = get_daily_summary(BD)
    assert summary["total_otomax"] == Decimal("0.00")
    assert summary["matched_auto_count"] == 1
    assert summary["unmatched_bank_amount"] == Decimal("20500.00")
    assert summary["selisih"] == Decimal("20500.00")


@pytest.mark.django_db
def test_range_summary_selisih_matches_daily_summary_definition():
    _bank("REF-COCOK-002", "100000", book_date=BD)
    _otomax("REF-COCOK-002", "100000", book_date=BD_NEXT)
    _bank("REF-TANPA-PASANGAN-2", "5000", book_date=BD)

    run_match(BD)

    summary = get_range_summary(BD, BD_NEXT)
    assert summary["selisih"] == Decimal("5000.00")


@pytest.mark.django_db
def test_per_bank_otomax_total_includes_all_categories_not_just_topup_tartun():
    """Rekapitulasi Per Bank sebelumnya cuma nampilin sisi bank -- sama sekali tidak ada
    total Otomax di situ. Sekarang 'otomax_total' per channel harus menjumlahkan SEMUA
    kategori Otomax (bukan cuma TOPUP_TARTUN) yang channel_hint-nya cocok."""
    _otomax("TARTUN TF BRI ABC", "500000", channel=Channel.BRI)
    _otomax_other("BAYAR KE BRI ABC", "50000", OtomaxCategory.PAYMENT, channel_hint=Channel.BRI)

    summary = get_daily_summary(BD)
    assert summary["per_bank"][Channel.BRI]["otomax_total"] == Decimal("550000.00")
    assert summary["per_bank"][Channel.BRI]["otomax_count"] == 2


@pytest.mark.django_db
def test_otomax_lain_lain_collects_entries_without_channel_hint():
    """Entri Otomax tanpa channel_hint (biaya admin, setor/ambil setoran) tidak masuk
    channel manapun -- harus tetap kelihatan, dikumpulkan di bucket lain_lain, bukan
    hilang begitu saja dari rekapitulasi."""
    _otomax_other("ADMIN BIAYA BULANAN", "-15000", OtomaxCategory.ADMIN, channel_hint="")
    _otomax_other("STOR TUNAI KE KAS", "300000", OtomaxCategory.STOR_IN, channel_hint="")
    _otomax("TARTUN TF BRI ABC", "500000", channel=Channel.BRI)  # kontrol: ini TIDAK boleh ikut kehitung

    summary = get_daily_summary(BD)
    assert summary["otomax_lain_lain"]["otomax_total"] == Decimal("285000.00")
    assert summary["otomax_lain_lain"]["otomax_count"] == 2


@pytest.mark.django_db
def test_otomax_lain_lain_excludes_ignored_entries():
    """Entri REVERSAL yang sudah dinetralkan (IGNORED) tidak boleh ikut dihitung sebagai
    saldo menggantung di bucket lain_lain -- itu koreksi internal yang sudah selesai."""
    _otomax_other(
        "REV SUDAH DINETRALKAN",
        "-40000",
        OtomaxCategory.REVERSAL,
        channel_hint="",
        match_status=MatchStatus.IGNORED,
    )

    summary = get_daily_summary(BD)
    assert summary["otomax_lain_lain"]["otomax_count"] == 0
    assert summary["otomax_lain_lain"]["otomax_total"] == Decimal("0.00")


@pytest.mark.django_db
def test_get_daily_rows_identik_dengan_get_daily_summary_per_tanggal():
    """Optimasi Laporan (1.381 -> ±6 query): rumus agregat per rentang HARUS menghasilkan
    angka yang sama persis dengan get_daily_summary() untuk setiap tanggal & kolom."""
    from apps.recon.models import ReconDay
    from apps.recon.reports import DAILY_ROW_KEYS, get_daily_rows
    from apps.recon.resolve import manual_pair_transactions, tag_manual_mutation

    d1, d2, d3 = BD, BD_NEXT, BD_NEXT + timedelta(days=1)
    # Hari 1: cocok otomatis (ref sama), sisa bank & Otomax tanpa pasangan, debit negatif.
    _bank("ATMLTRPRM 01884 000001039 21540100059656", "550000", book_date=d1)
    _otomax("TARTUN EDC BRI ATMLTRPRM 01884 000001039 21540100059656", "550000", book_date=d1)
    _bank("TRF MASUK TANPA PASANGAN 111", "125000", book_date=d1)
    _bank("DEBIT BIAYA", "-6500", book_date=d1)
    _otomax("TARTUN TF BRI TIDAK ADA DI BANK 222", "300000", book_date=d1)
    _otomax_other("STOR KAS", "90000", OtomaxCategory.OTHER, book_date=d1)
    run_match(d1)
    # Hari 2: cocok manual dengan nominal beda (AMOUNT_DIFF OPEN) + tag manual + REV netral.
    b_diff = _bank("QRIS CILENGKRANG 3 CELL 004769151", "2778000", channel=Channel.BCA, book_date=d2)
    o_diff = _otomax("TARTUN TF BCA CILENGKRANG 3", "2788000", channel=Channel.BCA, book_date=d2)
    manual_pair_transactions(b_diff, o_diff)
    # Selisih nominal kedua di hari yang sama: GROUP BY per tanggal harus menjumlahkan keduanya.
    b_diff2 = _bank("QRIS CIBIRU 7 CELL 004769999", "1500000", channel=Channel.BCA, book_date=d2)
    manual_pair_transactions(b_diff2, _otomax("TARTUN TF BCA CIBIRU 7", "1490000", channel=Channel.BCA, book_date=d2))
    tag_manual_mutation(_bank("SETOR TUNAI KASIR", "400000", book_date=d2), tag="setor_tunai")
    _otomax_other("TARTUN SALAH", "75000", OtomaxCategory.TOPUP_TARTUN, "BRI", d2, MatchStatus.IGNORED)
    _otomax("TARTUN TF BRI MASIH PENDING 333", "80000", book_date=d2)
    run_match(d2)
    # Hari 3: tutup buku; hari 4 (di luar data) kosong.
    _bank("MUTASI HARI TIGA 444", "10000", book_date=d3)
    run_match(d3)
    ReconDay.objects.update_or_create(book_date=d3, defaults={"locked": True, "status": "CLOSED"})

    start, end = BD - timedelta(days=1), d3 + timedelta(days=1)
    rows = get_daily_rows(start, end)

    assert [r["book_date"] for r in rows] == [end - timedelta(days=i) for i in range((end - start).days + 1)]
    for row in rows:
        expected = get_daily_summary(row["book_date"])
        for key in DAILY_ROW_KEYS:
            assert row[key] == expected[key], (row["book_date"], key, row[key], expected[key])
    # Sanity: data uji memang mengisi kolom-kolom yang dibandingkan.
    by_date = {r["book_date"]: r for r in rows}
    assert by_date[d1]["matched_auto_count"] >= 1
    assert by_date[d2]["amount_diff_count"] == 2
    assert by_date[d2]["matched_manual_count"] >= 1
    assert by_date[d3]["day_locked"] is True


@pytest.mark.django_db
def test_laporan_satu_bulan_tidak_lagi_ribuan_query(client):
    from django.contrib.auth import get_user_model
    from django.db import connection
    from django.test.utils import CaptureQueriesContext

    client.force_login(get_user_model().objects.create_user(username="op", password="password123"))
    with CaptureQueriesContext(connection) as q:
        res = client.get("/reports/", {"start_date": "2026-08-29", "end_date": "2026-09-28"})
    assert res.status_code == 200
    assert len(res.context["daily_rows"]) == 31
    assert len(q) < 80, len(q)  # dulu ±1.381 query untuk 31 hari


def _per_bank_breakdown_lama(bank_qs, otomax_all_qs):
    """Salinan PERSIS implementasi lama (7 query per bank) sebagai acuan pembanding."""
    from django.db.models import Sum

    from apps.core.enums import BANK_CHANNELS

    zero = Decimal("0.00")
    per_bank = {}
    for ch in BANK_CHANNELS:
        ch_mut = bank_qs.filter(channel=ch)
        ch_otomax = otomax_all_qs.filter(channel_hint=ch)
        per_bank[ch] = {
            "total": ch_mut.aggregate(s=Sum("amount"))["s"] or zero,
            "matched_auto": ch_mut.filter(match_status=MatchStatus.MATCHED).aggregate(s=Sum("amount"))["s"] or zero,
            "matched_manual": ch_mut.filter(match_status=MatchStatus.MANUAL).aggregate(s=Sum("amount"))["s"] or zero,
            "unmatched": ch_mut.filter(match_status=MatchStatus.UNMATCHED).aggregate(s=Sum("amount"))["s"] or zero,
            "unmatched_count": ch_mut.filter(match_status=MatchStatus.UNMATCHED).count(),
            "otomax_total": ch_otomax.aggregate(s=Sum("amount"))["s"] or zero,
            "otomax_count": ch_otomax.count(),
        }
    lain = otomax_all_qs.exclude(channel_hint__in=BANK_CHANNELS)
    return per_bank, {
        "otomax_total": lain.aggregate(s=Sum("amount"))["s"] or zero,
        "otomax_count": lain.count(),
    }


@pytest.mark.django_db
def test_per_bank_breakdown_identik_dengan_versi_lama():
    """Optimasi Dashboard (30 -> 3 query): rekap per bank harus sama persis dengan
    implementasi lama untuk berbagai status, channel, nominal negatif, dan entri Otomax
    tanpa channel / dinetralkan."""
    from apps.ingest.models import BankMutation
    from apps.recon.reports import _per_bank_breakdown
    from apps.recon.resolve import manual_pair_transactions, tag_manual_mutation

    _bank("ATMLTRPRM 01884 000001039 21540100059656", "550000")
    _otomax("TARTUN EDC BRI ATMLTRPRM 01884 000001039 21540100059656", "550000")
    _bank("BCA MASUK TANPA PASANGAN 1", "125000", channel=Channel.BCA)
    _bank("BCA MASUK TANPA PASANGAN 2", "75000", channel=Channel.BCA)
    _bank("DEBIT BIAYA", "-6500")
    b_diff = _bank("QRIS CILENGKRANG 3 CELL 004769151", "2778000", channel=Channel.MERCHANT_BCA)
    manual_pair_transactions(b_diff, _otomax("TARTUN QR CILENGKRANG", "2788000", channel=Channel.MERCHANT_BCA))
    tag_manual_mutation(_bank("SETOR TUNAI KASIR", "400000"), tag="setor_tunai")
    _otomax_other("ADMIN TARTUN", "-6500", OtomaxCategory.ADMIN)
    _otomax_other("STOR KAS", "90000", OtomaxCategory.OTHER, channel_hint="OTOMAX")
    _otomax_other("SALAH TEMBAK", "75000", OtomaxCategory.TOPUP_TARTUN, "BRI", BD, MatchStatus.IGNORED)
    _otomax("TARTUN TF MANDIRI PENDING", "80000", channel=Channel.MANDIRI)
    run_match(BD)

    bank_qs = BankMutation.objects.filter(book_date=BD)
    otomax_all_qs = OtomaxEntry.objects.filter(book_date=BD).exclude(match_status=MatchStatus.IGNORED)
    assert _per_bank_breakdown(bank_qs, otomax_all_qs) == _per_bank_breakdown_lama(bank_qs, otomax_all_qs)
    # Juga untuk rentang (dipakai get_range_summary) dan hari kosong.
    for qs_b, qs_o in [
        (BankMutation.objects.filter(book_date__range=(BD, BD_NEXT)), OtomaxEntry.objects.all()),
        (BankMutation.objects.filter(book_date=BD_NEXT), OtomaxEntry.objects.filter(book_date=BD_NEXT)),
    ]:
        assert _per_bank_breakdown(qs_b, qs_o) == _per_bank_breakdown_lama(qs_b, qs_o)


@pytest.mark.django_db
def test_generate_excel_discrepancies_and_export_view(client, django_user_model):
    """Test ekspor daftar selisih ke Excel (.xlsx) menghasilkan file valid dengan konten lengkap."""
    import io

    import openpyxl

    from apps.core.enums import DiscrepancyKind, DiscrepancyStatus
    from apps.recon.models import Discrepancy
    from apps.recon.reports import generate_excel_discrepancies

    user = django_user_model.objects.create_user(username="operator_export", password="password123")
    client.force_login(user)

    bm = _bank("TRANSFER DARI RESELLER XYZ", "1500000", channel=Channel.BCA, book_date=BD)
    oe = _otomax("TARTUN RESELLER XYZ", "1500000", channel=Channel.BCA, book_date=BD)

    Discrepancy.objects.create(
        code="SLS-20260911-001",
        origin_book_date=BD,
        channel=Channel.BCA,
        kind=DiscrepancyKind.BANK_ONLY,
        bank_mutation=bm,
        amount=Decimal("1500000.00"),
        status=DiscrepancyStatus.OPEN,
        note="Belum ada di Otomax",
    )
    Discrepancy.objects.create(
        code="SLS-20260911-002",
        origin_book_date=BD,
        channel=Channel.BCA,
        kind=DiscrepancyKind.OTOMAX_ONLY,
        otomax_entry=oe,
        amount=Decimal("-1500000.00"),
        status=DiscrepancyStatus.OPEN,
    )

    # 1. Test generate_excel_discrepancies service function
    excel_bytes = generate_excel_discrepancies(
        Discrepancy.objects.filter(origin_book_date=BD).order_by("code"),
        title_note="Tanggal Buku: 11 Sep 2026",
    )
    assert len(excel_bytes) > 0

    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    ws = wb.active
    assert ws.title == "Daftar Selisih"
    assert "DAFTAR SELISIH REKONSILIASI" in str(ws["A1"].value)
    assert "Tanggal Buku: 11 Sep 2026" in str(ws["A2"].value)

    # Verifikasi baris 5 dan 6 (data selisih)
    row5_vals = [ws.cell(row=5, column=col).value for col in range(1, 16)]
    assert "SLS-20260911-001" in row5_vals
    assert 1500000.0 in row5_vals

    row6_vals = [ws.cell(row=6, column=col).value for col in range(1, 16)]
    assert "SLS-20260911-002" in row6_vals
    assert -1500000.0 in row6_vals

    # 2. Test view endpoint /selisih/export/
    resp = client.get(f"/selisih/export/?d={BD.isoformat()}&status=OPEN")
    assert resp.status_code == 200
    assert resp["Content-Type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert 'attachment; filename="daftar_selisih_20260911.xlsx"' in resp["Content-Disposition"]

    wb_resp = openpyxl.load_workbook(io.BytesIO(resp.content))
    ws_resp = wb_resp.active
    assert ws_resp.title == "Daftar Selisih"


@pytest.mark.django_db
def test_generate_whatsapp_recon_text_and_note_update(client, django_user_model):
    """Test pembuatan teks laporan WhatsApp dan pembaruan catatan investigasi selisih."""
    from apps.core.enums import DiscrepancyKind, DiscrepancyStatus
    from apps.recon.models import Discrepancy
    from apps.recon.wa_report import generate_whatsapp_recon_text

    user = django_user_model.objects.create_user(username="operator_wa", password="password123")
    client.force_login(user)

    d_prev = date(2026, 9, 10)
    d_curr = date(2026, 9, 11)

    bm_prev = _bank("SETORAN TARTUN LAMA", "500000", channel=Channel.BRI, book_date=d_prev)
    bm_curr = _bank("AISUMIATI-BANK KESEJAHTERAAN", "700000", channel=Channel.BRI, book_date=d_curr)

    # Selisih hari sebelumnya (1 resolved, 1 open)
    Discrepancy.objects.create(
        code="SLS-20260910-001",
        origin_book_date=d_prev,
        channel=Channel.BRI,
        kind=DiscrepancyKind.BANK_ONLY,
        bank_mutation=bm_prev,
        amount=Decimal("500000.00"),
        status=DiscrepancyStatus.RESOLVED,
        resolved_book_date=d_curr,
    )
    Discrepancy.objects.create(
        code="SLS-20260910-002",
        origin_book_date=d_prev,
        channel=Channel.MANDIRI,
        kind=DiscrepancyKind.BANK_ONLY,
        amount=Decimal("100000.00"),
        status=DiscrepancyStatus.OPEN,
        note="Follow up konter",
    )

    # Selisih hari ini (1 open)
    disc_curr = Discrepancy.objects.create(
        code="SLS-20260911-001",
        origin_book_date=d_curr,
        channel=Channel.BRI,
        kind=DiscrepancyKind.BANK_ONLY,
        bank_mutation=bm_curr,
        amount=Decimal("700000.00"),
        status=DiscrepancyStatus.OPEN,
    )

    # 1. Test update note endpoint
    resp_note = client.post(
        f"/selisih/{disc_curr.pk}/note/",
        {"note": "Saldo masuk BRI belum ada claim", "next_url": "/selisih/?d=2026-09-11"},
        follow=True,
    )
    assert resp_note.status_code == 200
    disc_curr.refresh_from_db()
    assert disc_curr.note == "Saldo masuk BRI belum ada claim"

    # 2. Test generate_whatsapp_recon_text
    wa_text = generate_whatsapp_recon_text(d_curr)
    assert "LAPORAN REKONSILIASI HARIAN" in wa_text
    assert "11 September 2026" in wa_text
    assert "OUTSTANDING HARI INI" in wa_text
    assert "700.000" in wa_text
    assert "AISUMIATI-BANK KESEJAHTERAAN" in wa_text
    assert "Saldo masuk BRI belum ada claim" in wa_text
    assert "UPDATE PROGRESS TGL SEBELUMNYA (10 Sep 2026)" in wa_text
    assert "Sudah Selesai / Klop (1)" in wa_text
    assert "SLS-20260910-001" in wa_text
    assert "Masih Belum Selesai (1)" in wa_text
    assert "Follow up konter" in wa_text

    # 3. Test endpoint /selisih/wa-text/
    resp_wa = client.get(f"/selisih/wa-text/?d={d_curr.isoformat()}")
    assert resp_wa.status_code == 200
    json_data = resp_wa.json()
    assert json_data["status"] == "ok"
    assert "LAPORAN REKONSILIASI HARIAN" in json_data["text"]


