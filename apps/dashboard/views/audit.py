"""Halaman Audit Data — semua transaksi mentah (bank, debit, Otomax, dikecualikan) untuk satu
tanggal atau rentang, dengan filter detail, export Excel, dan hapus per baris mentah.

Mode tanggal:
- Per tanggal (`?d=`, dari sidebar / bilah atas).
- Rentang (`?start_date=&end_date=`), dibatasi audit.MAX_RANGE_DAYS hari.

Hapus: pilih baris -> panel ringkasan (periksa) -> ketik HAPUS. Yang dihapus adalah baris
mentahnya; pasangan cocok/netralnya dibatalkan dulu (lihat apps.recon.raw_delete).
"""

from __future__ import annotations

import io
from datetime import timedelta

import openpyxl
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.http import Http404, HttpResponse
from django.shortcuts import redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_POST
from openpyxl.styles import Alignment, Font, PatternFill

from apps.core.enums import Channel
from apps.recon import audit
from apps.recon.models import Match
from apps.recon.raw_delete import count_paired, delete_raw_rows, plan_delete

from ._shared import _paginate_rows, _parse_date

PAGE_SIZE = 100
MAX_DELETE = 3000  # lebih dari ini: pakai Reset & Purge Data (per hari), bukan hapus per baris
SOURCE_CHOICES = [(Channel.BRI, "BRI"), (Channel.BCA, "BCA"), (Channel.MANDIRI, "Mandiri")]
SOURCE_CHOICES += [(Channel.MERCHANT_BCA, "Merchant BCA (QRIS)"), (Channel.OTOMAX, "Otomax")]
SIDES = [("", "Semua"), (audit.SIDE_BANK, "Bank"), (audit.SIDE_OTOMAX, "Otomax")]


def _filters(params) -> tuple[audit.AuditFilters, bool]:
    """Filter dari query string / form POST. Return (filters, rentang_dipotong)."""
    if params.get("d") or not (params.get("start_date") or params.get("end_date")):
        start = end = _parse_date(params.get("d"))
    else:
        start = _parse_date(params.get("start_date"), default=_parse_date(params.get("end_date")))
        end = _parse_date(params.get("end_date"), default=start)
        if start > end:
            start, end = end, start
    clipped = (end - start).days + 1 > audit.MAX_RANGE_DAYS
    if clipped:
        end = start + timedelta(days=audit.MAX_RANGE_DAYS - 1)
    side = params.get("side", "")
    source = params.get("source", "")
    status = params.get("status", "")
    return (
        audit.AuditFilters(
            start=start,
            end=end,
            side=side if side in (audit.SIDE_BANK, audit.SIDE_OTOMAX) else "",
            source=source if source in Channel.values else "",
            status=status if status in audit.STATUS_LABELS else "",
            q=params.get("q", "").strip()[:100],
        ),
        clipped,
    )


def _summary_cards(result: audit.AuditResult, filters: audit.AuditFilters) -> list[dict]:
    cards = []
    for side, title in ((audit.SIDE_BANK, "Mutasi Bank"), (audit.SIDE_OTOMAX, "Otomax")):
        data = result.summary.get(side, {})
        items = [
            {
                "status": s,
                "label": audit.STATUS_LABELS[s],
                "tone": audit.STATUS_TONES[s],
                "active": filters.status == s and filters.side == side,
                **data[s],
            }
            for s in audit.STATUS_LABELS
            if s in data
        ]
        cards.append(
            {
                "side": side,
                "title": title,
                "items": items,
                "n": sum(i["n"] for i in items),
                "amount": sum((i["amount"] for i in items), 0),
            }
        )
    return cards


def _query_without(request, *keys) -> str:
    params = request.GET.copy()
    for k in (*keys, "page"):
        params.pop(k, None)
    return params.urlencode()


@login_required
def audit_data_view(request):
    if not any(request.GET.get(k) for k in ("d", "start_date", "end_date")):
        # Selalu punya tanggal eksplisit di URL, supaya bilah atas & isi halaman sama.
        d = request.session.get("last_book_date") or timezone.localdate().isoformat()
        params = request.GET.copy()
        params["d"] = d
        return redirect(f"{request.path}?{params.urlencode()}")

    filters, clipped = _filters(request.GET)
    result = audit.collect(filters)
    page_obj, is_more, list_url = _paginate_rows(request, result.rows, PAGE_SIZE)
    ctx = {"page_obj": page_obj, "list_url": list_url, "filters": filters}
    if is_more:
        return render(request, "dashboard/_audit_rows.html", ctx)

    if clipped:
        messages.warning(
            request,
            f"Rentang dibatasi {audit.MAX_RANGE_DAYS} hari: ditampilkan sampai {filters.end.strftime('%d %b %Y')}.",
        )
    range_mode = not request.GET.get("d")
    total_amount = {"bank": 0, "otomax": 0}
    for r in result.rows:
        total_amount[r.side] += r.amount
    ctx.update(
        {
            "book_date": None if range_mode else filters.start,
            "range_mode": range_mode,
            "result_count": len(result.rows),
            "total_amount": total_amount,
            "cards": _summary_cards(result, filters),
            "sides": SIDES,
            "sources": SOURCE_CHOICES,
            "statuses": list(audit.STATUS_LABELS.items()),
            "has_filter": bool(filters.side or filters.source or filters.status or filters.q),
            "base_query": _query_without(request, "side", "source", "status", "q"),
            "side_query": _query_without(request, "side"),
            "status_query": _query_without(request, "status", "side"),
            "export_query": _query_without(request),
            # Field filter aktif: ikut dikirim saat "pilih semua hasil filter" dihapus.
            "filter_items": [
                (k, request.GET[k])
                for k in ("d", "start_date", "end_date", "side", "source", "status", "q")
                if request.GET.get(k)
            ],
            "range_start": filters.start if range_mode else filters.start - timedelta(days=6),
            "max_range": audit.MAX_RANGE_DAYS,
            "max_delete": MAX_DELETE,
        }
    )
    return render(request, "dashboard/audit.html", ctx)


def _detail_fields(src, obj) -> list[tuple[str, object]]:
    b = obj.import_batch
    common_tail = [
        ("File asal", b.source_filename if b else "-"),
        ("Diupload", timezone.localtime(b.created_at).strftime("%d %b %Y %H:%M") if b and b.created_at else "-"),
        ("Oleh", (b.uploaded_by.get_username() if b and b.uploaded_by_id else "-")),
        ("ID baris", f"{src}:{obj.pk}"),
    ]
    when = getattr(obj, "txn_datetime", None) or getattr(obj, "entry_datetime", None)
    when_s = timezone.localtime(when).strftime("%d %b %Y %H:%M:%S") if when else "-"
    if src == audit.SRC_BANK:
        return [
            ("Channel", obj.get_channel_display()),
            ("Tanggal buku", obj.book_date),
            ("Waktu transaksi", when_s),
            ("Keterangan", obj.description_raw),
            ("Outlet", obj.outlet_name or "-"),
            ("No. referensi", obj.external_ref or "-"),
            ("Ref inti", obj.ref_core or "-"),
            ("Token", ", ".join(obj.extracted_tokens or []) or "-"),
            ("Status mesin", obj.get_match_status_display()),
            ("Tag manual", obj.get_tag_manual_display() if obj.tag_manual else "-"),
            ("Catatan", obj.manual_note or "-"),
            *common_tail,
        ]
    if src == audit.SRC_OTOMAX:
        return [
            ("Tanggal buku", obj.book_date),
            ("Waktu entri", when_s),
            ("Reseller", obj.reseller_name_raw),
            ("Keterangan", obj.description_raw),
            ("Kategori", obj.get_category_display()),
            ("Petunjuk bank", obj.channel_hint or "-"),
            ("Ref inti", obj.ref_core or "-"),
            ("Token", ", ".join(obj.extracted_tokens or []) or "-"),
            ("Status mesin", obj.get_match_status_display()),
            ("Lawan netral", f"#{obj.net_pair_id}" if obj.net_pair_id else "-"),
            ("Catatan", obj.note or "-"),
            *common_tail,
        ]
    if src == audit.SRC_DEBIT:
        return [
            ("Channel", obj.get_channel_display()),
            ("Tanggal buku", obj.book_date),
            ("Waktu transaksi", when_s),
            ("Keterangan", obj.description_raw),
            *common_tail,
        ]
    return [
        ("Sisi", "Bank" if obj.source_type == "BANK" else "Otomax"),
        ("Channel", obj.channel),
        ("Tanggal buku", obj.book_date),
        ("Waktu", when_s),
        ("Keterangan", obj.description_raw),
        ("Kategori", obj.category or "-"),
        ("Alasan", obj.reason or "-"),
        *common_tail,
    ]


@login_required
def audit_detail_view(request, src: str, pk: int):
    """Panel detail satu baris mentah: semua kolom + riwayat pencocokan + selisihnya."""
    model = audit.SOURCE_MODELS.get(src)
    if model is None:
        raise Http404
    obj = model.objects.select_related("import_batch", "import_batch__uploaded_by").filter(pk=pk).first()
    if obj is None:
        return HttpResponse(
            '<div class="callout callout-warn text-xs">Baris ini sudah tidak ada (mungkin baru dihapus).</div>'
        )
    matches, discrepancies = [], []
    if src in (audit.SRC_BANK, audit.SRC_OTOMAX):
        fk, m2m = ("bank_mutation", "bank_mutations") if src == audit.SRC_BANK else ("otomax_entry", "otomax_entries")
        matches = list(
            (Match.objects.filter(**{fk: obj}) | Match.objects.filter(**{m2m: obj}))
            .distinct()
            .select_related("bank_mutation", "otomax_entry", "voided_by")
            .order_by("-created_at")[:10]
        )
        discrepancies = list(obj.discrepancies.order_by("-created_at")[:10])
    return render(
        request,
        "dashboard/_audit_detail.html",
        {"fields": _detail_fields(src, obj), "matches": matches, "discrepancies": discrepancies},
    )


def _selected_keys(request) -> list[str]:
    """Kunci baris dari form: daftar eksplisit, atau 'semua hasil filter' dihitung ulang."""
    if request.POST.get("all") == "1":
        filters, _ = _filters(request.POST)
        return [r.key for r in audit.collect(filters).rows]
    raw = request.POST.get("keys", "")
    return [k for k in raw.replace("\n", ",").split(",") if k.strip()]


@login_required
@require_POST
def audit_delete_preview(request):
    """Panel konfirmasi (HTMX): ringkasan apa yang akan dihapus. Tidak menulis apa pun."""
    keys = _selected_keys(request)
    too_many = len(keys) > MAX_DELETE
    plan = plan_delete(keys) if keys and not too_many else None
    return render(
        request,
        "dashboard/_audit_delete_confirm.html",
        {
            "plan": plan,
            "paired": count_paired(plan) if plan else 0,
            "keys": ",".join(f"{s}:{o.pk}" for s, o in plan.objects) if plan else "",
            "requested": len(keys),
            "too_many": too_many,
            "max_delete": MAX_DELETE,
            "next_url": _safe_next(request.POST.get("next", "")),
        },
    )


def _safe_next(url: str) -> str:
    base = reverse("audit-data")
    return url if url.startswith(base) else base


@login_required
@require_POST
def audit_delete_action(request):
    next_url = _safe_next(request.POST.get("next", ""))
    if request.POST.get("confirm", "").strip() != "HAPUS":
        messages.error(request, "Penghapusan dibatalkan: ketik HAPUS (huruf besar) untuk konfirmasi.")
        return redirect(next_url)
    keys = [k for k in request.POST.get("keys", "").split(",") if k.strip()]
    if not keys:
        messages.warning(request, "Tidak ada baris yang dipilih.")
        return redirect(next_url)
    if len(keys) > MAX_DELETE:
        messages.error(request, f"Maksimal {MAX_DELETE} baris sekali hapus.")
        return redirect(next_url)

    res = delete_raw_rows(keys, user=request.user)
    parts = [f"{res['deleted']} baris dihapus"]
    if res["unpaired"] or res["unnetted"]:
        parts.append(f"{res['unpaired'] + res['unnetted']} pasangan dibatalkan (lawannya kembali ke antrean)")
    if res["blocked"]:
        parts.append(f"{res['blocked']} ditolak karena tanggalnya sudah tutup buku")
    if res["missing"]:
        parts.append(f"{res['missing']} sudah tidak ada")
    (messages.success if res["deleted"] else messages.warning)(request, "; ".join(parts) + ".")
    return redirect(next_url)


@login_required
def audit_export_action(request):
    filters, _ = _filters(request.GET)
    rows = audit.collect(filters).rows

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Audit Data"
    period = (
        filters.start.strftime("%d %b %Y")
        if not filters.is_range
        else f"{filters.start.strftime('%d %b %Y')} s/d {filters.end.strftime('%d %b %Y')}"
    )
    ws["A1"] = "AUDIT DATA — SEMUA TRANSAKSI"
    ws["A1"].font = Font(name="Arial", size=13, bold=True)
    notes = [f"Periode: {period}"]
    if filters.side:
        notes.append(f"Sisi: {filters.side}")
    if filters.source:
        notes.append(f"Sumber: {filters.source}")
    if filters.status:
        notes.append(f"Status: {audit.STATUS_LABELS[filters.status]}")
    if filters.q:
        notes.append(f"Cari: {filters.q}")
    ws["A2"] = " · ".join(notes)

    headers = [
        "Tanggal Buku",
        "Waktu",
        "Sisi",
        "Sumber",
        "Pihak",
        "Keterangan",
        "Kategori",
        "Nominal",
        "Status",
        "Alasan",
        "File Asal",
        "Diupload",
        "ID Baris",
    ]
    fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=4, column=i, value=h)
        c.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        c.fill = fill
        c.alignment = Alignment(horizontal="center")

    for r_idx, r in enumerate(rows, start=5):
        when = timezone.localtime(r.when).strftime("%H:%M:%S") if r.when else ""
        up = timezone.localtime(r.uploaded_at).strftime("%Y-%m-%d %H:%M") if r.uploaded_at else ""
        values = [
            r.book_date,
            when,
            "Bank" if r.side == audit.SIDE_BANK else "Otomax",
            r.channel_label,
            r.party,
            r.description,
            r.category,
            r.amount,  # Decimal langsung, tanpa float
            r.status_label,
            r.reason,
            r.filename,
            up,
            r.key,
        ]
        for c_idx, v in enumerate(values, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=v)
            if c_idx == 1:
                cell.number_format = "yyyy-mm-dd"
            elif c_idx == 8:
                cell.number_format = "#,##0;[Red]-#,##0"
    widths = [12, 10, 8, 20, 22, 60, 16, 16, 24, 60, 30, 17, 10]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:{openpyxl.utils.get_column_letter(len(headers))}{max(4, len(rows) + 4)}"

    buf = io.BytesIO()
    wb.save(buf)
    name = f"audit_{filters.start}" + (f"_{filters.end}" if filters.is_range else "") + ".xlsx"
    response = HttpResponse(
        buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="{name}"'
    return response
