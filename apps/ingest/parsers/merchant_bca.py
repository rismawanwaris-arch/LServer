"""Parser rekap QRIS Merchant BCA — Mendukung Excel (.xlsx) sheet SUMMARY & TSV."""

from __future__ import annotations

import io
import re
from datetime import datetime
from decimal import Decimal, InvalidOperation

import openpyxl

from ._dates import parse_id_datetime
from .base import ParsedBankRow, ParseResult, to_text


def _parse_xlsx(file_bytes: bytes, result: ParseResult) -> bool:
    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    except Exception as exc:
        result.warnings.append(f"Gagal membaca format xlsx: {exc}")
        return False

    summary_sheet = None
    for name in wb.sheetnames:
        if "SUMMARY" in name.upper():
            summary_sheet = wb[name]
            break
    if summary_sheet is None:
        summary_sheet = wb.active

    rows = list(summary_sheet.iter_rows(values_only=True))
    header_idx = -1
    for idx, r in enumerate(rows[:10]):
        if not r:
            continue
        cells = [str(c).strip().upper() for c in r if c is not None]
        if any("MERCHANT NAME" in c for c in cells) and any("TOTAL" in c or "AMOUNT" in c for c in cells):
            header_idx = idx
            break

    start_idx = header_idx + 1 if header_idx != -1 else 4
    if start_idx >= len(rows):
        return False

    sheet_date: datetime | None = None
    for r in rows[:start_idx]:
        if not r:
            continue
        for c in r:
            if c:
                parsed_dt = parse_id_datetime(str(c))
                if parsed_dt:
                    sheet_date = parsed_dt
                    result.book_date = parsed_dt.date()
                    break
        if sheet_date:
            break

    name_col, mid_col, amt_col, freq_col = 0, 1, 3, 2
    if header_idx != -1:
        h_row = [str(c).strip().upper() if c is not None else "" for c in rows[header_idx]]
        for c_idx, val in enumerate(h_row):
            if "MERCHANT NAME" in val or "NAMA" in val:
                name_col = c_idx
            elif "MERCHANT ID" in val or "MID" in val:
                mid_col = c_idx
            elif "TOTAL AMOUNT" in val or "AMOUNT" in val or "JUMLAH" in val:
                amt_col = c_idx
            elif "FREQUENCY" in val or "FREK" in val:
                freq_col = c_idx

    for row in rows[start_idx:]:
        if not row or len(row) <= max(name_col, mid_col, amt_col):
            continue
        name_val = str(row[name_col] or "").strip()
        mid_val = str(row[mid_col] or "").strip()
        amt_val = row[amt_col]

        if not name_val or name_val.upper() in {"TOTAL", "GRAND TOTAL", "MERCHANT NAME"}:
            continue

        try:
            if isinstance(amt_val, int | float | Decimal):
                amount = Decimal(str(amt_val))
            else:
                clean_amt = re.sub(r"[^\d.]", "", str(amt_val or "0").replace(",", ""))
                amount = Decimal(clean_amt or "0")
        except (InvalidOperation, ValueError):
            continue

        freq = 0
        if len(row) > freq_col and row[freq_col] is not None:
            try:
                freq = int(re.sub(r"[^\d]", "", str(row[freq_col])) or "0")
            except ValueError:
                freq = 0

        mid = re.sub(r"[^\w]", "", mid_val)

        result.bank_rows.append(
            ParsedBankRow(
                description_raw=f"QRIS {name_val} {mid}",
                amount=amount,
                txn_datetime=sheet_date,
                external_ref=mid,
                outlet_name=name_val,
                frequency=freq,
            )
        )
    return len(result.bank_rows) > 0


def parse(data: str | bytes) -> ParseResult:
    result = ParseResult()

    if isinstance(data, bytes) and data.startswith(b"PK"):
        if _parse_xlsx(data, result):
            return result

    text = to_text(data)
    for line in text.splitlines():
        parts = [p.strip() for p in (line.split("\t") if "\t" in line else line.split(","))]
        if len(parts) < 3:
            continue
        name, mid = parts[0], parts[1]
        amount_raw = parts[-1]
        freq_raw = parts[2] if len(parts) >= 4 else "0"

        if name.lower() in {"merchant name", "total", "grand total"} or not re.search(r"\d{4,}", mid):
            continue
        try:
            amount = Decimal(re.sub(r"[^\d]", "", amount_raw))
            freq = int(re.sub(r"[^\d]", "", freq_raw) or 0)
        except (InvalidOperation, ValueError):
            result.warnings.append(f"baris merchant dilewati: {line!r}")
            continue

        clean_mid = re.sub(r"[^\w]", "", mid)
        result.bank_rows.append(
            ParsedBankRow(
                description_raw=f"QRIS {name} {clean_mid}",
                amount=amount,
                external_ref=clean_mid,
                outlet_name=name,
                frequency=freq,
            )
        )
    if not result.bank_rows:
        result.warnings.append("Tidak ada merchant terbaca — periksa sheet SUMMARY file Excel atau format TSV/CSV.")
    return result
